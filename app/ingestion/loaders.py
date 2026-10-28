"""Text and metadata extraction for the required file formats.

Supported uploads: PDF, DOCX, HTML, Markdown, plain text, JSON, and JSONL.
That is more than the three-format minimum. Metadata keeps the source
identity that later citations point back to.
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path

from app.textutil import normalize_text


@dataclass
class LoadedDocument:
    title: str
    text: str
    fmt: str
    metadata: dict = field(default_factory=dict)


class UnsupportedFormat(ValueError):
    pass


def load_path(path: Path) -> list[LoadedDocument]:
    suffix = path.suffix.lower()
    data = path.read_bytes()
    return load_bytes(path.name, data, suffix)


def load_bytes(filename: str, data: bytes, suffix: str | None = None) -> list[LoadedDocument]:
    suffix = (suffix or Path(filename).suffix).lower()
    if suffix in {".txt", ".md", ".markdown"}:
        return [_text_document(filename, data, "markdown" if suffix != ".txt" else "txt")]
    if suffix in {".html", ".htm"}:
        return [_html_document(filename, data)]
    if suffix == ".pdf":
        return [_pdf_document(filename, data)]
    if suffix == ".docx":
        return [_docx_document(filename, data)]
    if suffix == ".json":
        return _json_documents(filename, data)
    if suffix == ".jsonl":
        return _jsonl_documents(filename, data)
    raise UnsupportedFormat(
        f"Unsupported file type '{suffix}'. Use PDF, DOCX, HTML, Markdown, TXT, JSON, or JSONL."
    )


def _decode(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _text_document(filename: str, data: bytes, fmt: str) -> LoadedDocument:
    text = normalize_text(_decode(data))
    title = _title_from_text(text) or Path(filename).stem
    return LoadedDocument(
        title=title,
        text=text,
        fmt=fmt,
        metadata={"filename": filename, "bytes": len(data)},
    )


def _title_from_text(text: str) -> str:
    for line in text.splitlines():
        stripped = line.strip().lstrip("#").strip()
        if not stripped or re.fullmatch(r"\[page \d+\]", stripped):
            continue
        return stripped[:180]
    return ""


def _html_document(filename: str, data: bytes) -> LoadedDocument:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_decode(data), "html.parser")
    title_tag = soup.find("title")
    title = title_tag.get_text(" ", strip=True) if title_tag else Path(filename).stem
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    text = normalize_text(soup.get_text("\n"))
    return LoadedDocument(
        title=title or Path(filename).stem,
        text=text,
        fmt="html",
        metadata={"filename": filename, "bytes": len(data)},
    )


def _pdf_document(filename: str, data: bytes) -> LoadedDocument:
    from pypdf import PdfReader
    import io

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        extracted = page.extract_text() or ""
        if extracted.strip():
            pages.append(f"[page {index}]\n{extracted.strip()}")
    text = normalize_text("\n\n".join(pages))
    meta = reader.metadata
    title = ""
    author = ""
    if meta is not None:
        title = str(meta.title or "")
        author = str(meta.author or "")
    if not title or title == "None":
        title = _title_from_text(text) or Path(filename).stem
    return LoadedDocument(
        title=title[:180],
        text=text,
        fmt="pdf",
        metadata={
            "filename": filename,
            "bytes": len(data),
            "pages": len(reader.pages),
            "author": author,
        },
    )


def _docx_document(filename: str, data: bytes) -> LoadedDocument:
    import io

    import docx

    document = docx.Document(io.BytesIO(data))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    text = normalize_text("\n".join(paragraphs))
    core = document.core_properties
    title = (core.title or "").strip() or _title_from_text(text) or Path(filename).stem
    return LoadedDocument(
        title=title[:180],
        text=text,
        fmt="docx",
        metadata={
            "filename": filename,
            "bytes": len(data),
            "author": core.author or "",
            "paragraphs": len(paragraphs),
        },
    )


def _records_from_json(payload) -> list[dict]:
    if isinstance(payload, list):
        return [item for item in payload if isinstance(item, dict)]
    if isinstance(payload, dict):
        for key in ("documents", "records", "data"):
            if isinstance(payload.get(key), list):
                return [item for item in payload[key] if isinstance(item, dict)]
        return [payload]
    return []


def _document_from_record(filename: str, record: dict, fmt: str, index: int) -> LoadedDocument | None:
    text = record.get("text") or record.get("body") or record.get("content") or ""
    if not isinstance(text, str) or not text.strip():
        return None
    title = record.get("title") or record.get("name") or f"{Path(filename).stem} #{index + 1}"
    metadata = {
        "filename": filename,
        "record_index": index,
    }
    for key in ("source", "url", "author", "id"):
        if key in record and isinstance(record[key], (str, int, float)):
            metadata[key] = record[key]
    return LoadedDocument(
        title=str(title)[:180],
        text=normalize_text(text),
        fmt=fmt,
        metadata=metadata,
    )


def _json_documents(filename: str, data: bytes) -> list[LoadedDocument]:
    payload = json.loads(_decode(data))
    documents = []
    for index, record in enumerate(_records_from_json(payload)):
        loaded = _document_from_record(filename, record, "json", index)
        if loaded is not None:
            documents.append(loaded)
    if not documents:
        raise ValueError("JSON file did not contain any records with a text field.")
    return documents


def _jsonl_documents(filename: str, data: bytes) -> list[LoadedDocument]:
    documents = []
    for index, line in enumerate(_decode(data).splitlines()):
        if not line.strip():
            continue
        record = json.loads(line)
        if not isinstance(record, dict):
            continue
        loaded = _document_from_record(filename, record, "jsonl", index)
        if loaded is not None:
            documents.append(loaded)
    if not documents:
        raise ValueError("JSONL file did not contain any records with a text field.")
    return documents

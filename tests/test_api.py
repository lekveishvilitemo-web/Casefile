from fastapi.testclient import TestClient

import app.main as main
from app.ingestion.loaders import load_bytes


LEDGER = """Ledger service change log, March 2024.
On 2024-03-02 at 22:14 UTC the payments platform deployed ledger-service v2.4.1.
The change reduced the payment retry count from 5 to 1 and cut the database connection pool from 40 to 8.
"""


def test_upload_investigate_duplicate_and_delete(engine):
    main.engine = engine
    client = TestClient(main.app)

    created = client.post(
        "/api/documents/text",
        json={"title": "Change log", "text": LEDGER, "source": "seed:test"},
    )
    assert created.status_code == 200
    body = created.json()
    assert body["created"] is True
    doc_id = body["document"]["id"]

    duplicate = client.post(
        "/api/documents/text",
        json={"title": "Change log copy", "text": LEDGER, "source": "seed:test"},
    )
    assert duplicate.status_code == 200
    assert duplicate.json()["duplicate"] is True

    listed = client.get("/api/documents")
    assert listed.status_code == 200
    assert listed.json()["total"] == 1

    answer = client.post(
        "/api/investigate",
        json={"question": "What did ledger-service v2.4.1 do to the retry count?", "retrieval": "hybrid"},
    )
    assert answer.status_code == 200
    payload = answer.json()
    assert payload["abstained"] is False
    assert payload["citations"]
    assert any(step["tool"] == "search_evidence" for step in payload["steps"])

    missing = client.post(
        "/api/investigate",
        json={"question": "What color was the Zorblandia treasury flag in 1888?"},
    )
    assert missing.status_code == 200
    assert missing.json()["abstained"] is True
    assert missing.json()["citations"] == []

    replaced = client.put(
        f"/api/documents/{doc_id}",
        files={"file": ("note.txt", b"Replacement note about merchant-notifier only.", "text/plain")},
    )
    assert replaced.status_code == 200
    fetched = client.get(f"/api/documents/{doc_id}")
    assert "merchant-notifier" in fetched.json()["text"]

    deleted = client.delete(f"/api/documents/{doc_id}")
    assert deleted.status_code == 200
    assert client.get(f"/api/documents/{doc_id}").status_code == 404


def test_loaders_cover_text_html_and_json():
    text = load_bytes("note.txt", b"Alpha ledger note\nThe pool size is 8.")
    assert text[0].fmt == "txt"
    html = load_bytes("page.html", b"<html><title>Postmortem</title><body><p>Rollback at 11:40.</p></body></html>")
    assert html[0].title == "Postmortem"
    assert "Rollback" in html[0].text
    records = load_bytes(
        "rows.jsonl",
        b'{"title":"One","text":"First record about payments."}\n{"title":"Two","text":"Second record."}\n',
    )
    assert len(records) == 2


def test_pdf_and_docx_loaders():
    from io import BytesIO

    import docx
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(0, 8, "SRE timeline. Rollback of ledger-service completed at 11:40 UTC.")
    pdf_bytes = bytes(pdf.output())
    loaded_pdf = load_bytes("timeline.pdf", pdf_bytes)
    assert loaded_pdf[0].fmt == "pdf"
    assert "Rollback" in loaded_pdf[0].text

    document = docx.Document()
    document.add_paragraph("Vendor status memo")
    document.add_paragraph("The March payment delay was caused by an OrbitalNet regional outage.")
    buffer = BytesIO()
    document.save(buffer)
    loaded_docx = load_bytes("memo.docx", buffer.getvalue())
    assert loaded_docx[0].fmt == "docx"
    assert "OrbitalNet" in loaded_docx[0].text
